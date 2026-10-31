"""Reading files: every supported format, encodings, delimiters, and clear errors."""

from __future__ import annotations

import json
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from data_storyteller.ingest import IngestError, decode, read_source, sniff_delimiter


def test_delimited_text_is_read_as_text(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("code,amount\n007,NA\n010,1.50\n", encoding="utf-8")
    table = read_source(path).tables["data"]
    assert table.header is None
    assert table.rows == [["code", "amount"], ["007", "NA"], ["010", "1.50"]]  # nothing guessed


@pytest.mark.parametrize(
    ("text", "delimiter"),
    [
        ("a,b,c\n1,2,3\n", ","),
        ("a;b;c\n1,5;2;3\n", ";"),
        ("a\tb\n1\t2\n", "\t"),
        ("a|b\n1|2\n", "|"),
    ],
)
def test_delimiters(text: str, delimiter: str) -> None:
    assert sniff_delimiter(text) == delimiter


def test_encodings() -> None:
    assert decode("é".encode("cp1252")) == ("é", "cp1252", False)
    assert decode(b"\xef\xbb\xbfa") == ("a", "utf-8-sig", True)
    assert decode("é".encode())[1] == "utf-8"
    assert decode(b"\x81")[1] == "latin-1"  # not valid in Windows-1252 either


def test_a_non_utf8_file_is_noted(messy: dict[str, Path]) -> None:
    source = read_source(messy["windows1252"])
    assert source.encoding == "cp1252" and "Windows-1252" in source.notes[0]


def test_excel_keeps_types_and_every_sheet(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Orders"
    first.append(["date", "amount"])
    first.append([datetime(2024, 1, 2), 10.5])
    workbook.create_sheet("Empty")
    workbook.create_sheet("Notes").append(["just a note"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    source = read_source(path)
    assert list(source.tables) == ["Orders", "Notes"]  # the empty sheet is skipped
    assert source.tables["Orders"].rows[1] == [datetime(2024, 1, 2), 10.5]


def test_typed_formats_know_their_header(tmp_path: Path) -> None:
    frame = pd.DataFrame(
        {"id": [1, 2], "when": pd.to_datetime(pd.Series(["2024-01-01", None])), "x": [1.5, None]}
    )
    frame.to_parquet(tmp_path / "t.parquet")
    table = read_source(tmp_path / "t.parquet").tables["t"]
    assert table.header == ["id", "when", "x"]
    assert table.rows == [[1, datetime(2024, 1, 1), 1.5], [2, None, None]]

    (tmp_path / "t.json").write_text(
        json.dumps({"data": [{"a": 1, "b": {"c": "x"}}]}), encoding="utf-8"
    )
    assert read_source(tmp_path / "t.json").tables["t"].header == ["a", "b.c"]


def test_sqlite_is_opened_read_only(tmp_path: Path) -> None:
    path = tmp_path / "shop.db"
    with sqlite3.connect(path) as connection:
        connection.execute('CREATE TABLE "my orders" (id INTEGER, total REAL)')
        connection.execute('INSERT INTO "my orders" VALUES (1, 9.5)')
    connection.close()
    source = read_source(path)
    assert source.tables["my orders"].rows == [[1, 9.5]]
    with sqlite3.connect(path) as connection:
        assert connection.execute('SELECT COUNT(*) FROM "my orders"').fetchone() == (1,)
    connection.close()


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("missing.csv", None, "No file"),
        ("old.xls", b"x", "save it as .xlsx"),
        ("data.xml", b"<a/>", "Unsupported file type"),
        ("empty.csv", b"\n\n", "is empty"),
        ("list.json", b"[1, 2]", "not a table"),
    ],
)
def test_clear_errors(tmp_path: Path, name: str, content: bytes | None, message: str) -> None:
    path = tmp_path / name
    if content is not None:
        path.write_bytes(content)
    with pytest.raises(IngestError, match=message):
        read_source(path)


def test_the_upload_picker_offers_exactly_the_types_the_reader_supports(tmp_path: Path) -> None:
    from data_storyteller.models import ACCEPTED

    for suffix in ACCEPTED:  # every offered type is read, not refused as unsupported
        path = tmp_path / f"empty{suffix}"
        path.write_bytes(b"")
        with pytest.raises(IngestError) as caught:
            read_source(str(path))
        assert "Unsupported file type" not in str(caught.value), suffix
    unsupported = tmp_path / "data.xyz"
    unsupported.write_text("a")
    with pytest.raises(IngestError, match="Unsupported file type"):
        read_source(str(unsupported))
